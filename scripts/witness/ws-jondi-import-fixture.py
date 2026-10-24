#!/usr/bin/env python3
"""Build a fictional, disposable DOCX for manuscript-import regression tests."""
from pathlib import Path
import sys
from docx import Document

out = Path(sys.argv[1] if len(sys.argv) > 1 else '/tmp/ws-jondi-import-fixture-20261011.docx')
if out.exists():
    print('REUSING_EXISTING_TEST_DOCX', out)
    sys.exit(0)
doc = Document()
doc.add_heading('The Temporary Book', 0)
for chapter in range(1, 7):
    doc.add_heading(f'Chapter {chapter}: Return and Recognition', 1)
    for section in range(1, 31):
        doc.add_heading(f'Section {chapter}.{section}: A Small Passage', 2)
        doc.add_paragraph(
            f'The writer returns to a Work in chapter {chapter}, passage {section}. '
            'It is still the same Work, holding the same thoughts and decisions. '
            'The words are synthetic and belong only to an automated test. '
            'An import must preserve each passage, and an editorial return must recover it in context.'
        )
doc.save(out)
print('DISPOSABLE_DOCX_READY', out, out.stat().st_size)
